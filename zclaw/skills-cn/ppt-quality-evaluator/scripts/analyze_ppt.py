#!/usr/bin/env python3
"""
PPT质量评估脚本 - 分析PPTX文件的多维度质量指标
输出详细评分报告
"""
import sys, os, json
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.enum.text import PP_ALIGN
from collections import Counter
from math import sqrt

def analyze_pptx(path):
    """分析PPTX文件的多个质量维度"""
    if not os.path.exists(path):
        return {"error": f"文件不存在: {path}"}
    
    prs = Presentation(path)
    slides = prs.slides
    n = len(slides)
    sw, sh = prs.slide_width, prs.slide_height
    
    report = {
        "file": os.path.basename(path),
        "slide_count": n,
        "slide_size": f"{sw/914400:.1f}x{sh/914400:.1f}in",
        "aspect": f"{sw/sh:.2f}:1",
        "dimensions": {}
    }
    
    # ── 统计容器 ──
    all_fonts = []
    all_font_sizes = []
    all_colors = []
    shape_count = 0
    text_count = 0
    chart_count = 0
    image_count = 0
    overlap_count = 0
    overflow_count = 0
    empty_count = 0
    has_header_footer = [False, False]  # [header, footer]
    margin_violations = 0
    total_text_len = 0
    alignments = Counter()
    shapes_per_slide = []
    
    # 标准边距（英寸）
    STD_MARGIN = 0.4 * 914400  # 0.4in in EMU
    
    for idx, slide in enumerate(slides):
        slide_shapes = list(slide.shapes)
        shape_count += len(slide_shapes)
        slide_text_len = 0
        shape_positions = []  # ← 每页独立重置，防止跨页误判重叠
        
        # 检查是否有顶部底部条
        for shape in slide_shapes:
            l, t, w, h = shape.left, shape.top, shape.width, shape.height
            
            # 检查顶部装饰条 (thin, full-width bar at top)
            if t < Emu(10000) and h < Emu(100000) and w > 0.8*sw:
                has_header_footer[0] = True
            # 检查底部栏
            if t + h > 0.92*sh and h < Emu(250000) and w > 0.8*sw:
                has_header_footer[1] = True
            
            # 边距检查（非全宽形状应在边距内）
            if w < 0.9*sw and l > Emu(50000):
                if l < STD_MARGIN - Emu(50000):
                    margin_violations += 1
            
            # 颜色提取
            try:
                if hasattr(shape, 'fill') and shape.fill.type is not None:
                    fc = shape.fill.fore_color
                    if fc and fc.type is not None:
                        try:
                            all_colors.append(str(fc.rgb))
                        except: pass
            except: pass
            
            try:
                if hasattr(shape, 'line') and shape.line.fill.type is not None:
                    lc = shape.line.color
                    if lc and lc.type is not None:
                        try:
                            all_colors.append(str(lc.rgb))
                        except: pass
            except: pass
            
            # 文本分析
            if shape.has_text_frame:
                text_count += 1
                for para in shape.text_frame.paragraphs:
                    txt = para.text.strip()
                    if txt:
                        slide_text_len += len(txt)
                        total_text_len += len(txt)
                        if para.font.size:
                            all_font_sizes.append(para.font.size)
                        if para.font.name:
                            all_fonts.append(para.font.name)
                        if para.alignment:
                            alignments[para.alignment] += 1
            
            # 图片
            if shape.shape_type == 13:  # Picture
                image_count += 1
            
            # 图表
            if hasattr(shape, 'chart'):
                chart_count += 1
            
            # 记录位置
            shape_positions.append((l, t, w, h))
        
        # 空页检查
        if slide_text_len < 20:
            empty_count += 1
        
        # 记录每页形状数（用于视觉一致性评估）
        shapes_per_slide.append(len(slide_shapes))
        
        # ── 本页重叠检查（每页独立检测，防止跨页误判） ──
        for i in range(len(shape_positions)):
            for j in range(i+1, len(shape_positions)):
                l1, t1, w1, h1 = shape_positions[i]
                l2, t2, w2, h2 = shape_positions[j]
                # 忽略微小形状（装饰线、小标签等）
                if w1 < Emu(80000) or h1 < Emu(80000) or w2 < Emu(80000) or h2 < Emu(80000):
                    continue
                # 忽略极扁的形状（色条、分隔线）
                if h1 < Emu(60000) or h2 < Emu(60000):
                    continue
                # 忽略全宽背景形状
                if w1 > 0.9*sw or w2 > 0.9*sw:
                    continue
                # 矩形重叠检测
                if l1 < l2+w2 and l1+w1 > l2 and t1 < t2+h2 and t1+h1 > t2:
                    # 检查是否是一个形状完全包含另一个（卡片内的装饰条属于设计模式，不计为重叠）
                    contained = (l1 >= l2 and t1 >= t2 and l1+w1 <= l2+w2 and t1+h1 <= t2+h2) or \
                                (l2 >= l1 and t2 >= t1 and l2+w2 <= l1+w1 and t2+h2 <= t1+h1)
                    if contained:
                        continue
                    # 计算重叠面积占比
                    ox = max(0, min(l1+w1, l2+w2) - max(l1, l2))
                    oy = max(0, min(t1+h1, t2+h2) - max(t1, t2))
                    overlap_area = ox * oy
                    min_area = min(w1*h1, w2*h2)
                    if min_area > 0 and overlap_area / min_area > 0.15:
                        overlap_count += 1
    
    # ════════════════════════════════════
    # 维度评分
    # ════════════════════════════════════
    
    scores = {}
    
    # 1. 页面结构 (15分)
    struct_score = 15
    struct_notes = []
    if n < 5: struct_score -= 3; struct_notes.append("页数偏少(<5)")
    if n > 50: struct_score -= 2; struct_notes.append("页数过多(>50)")
    if has_header_footer[0]: struct_notes.append("✓ 有顶部装饰条")
    else: struct_score -= 2; struct_notes.append("✗ 缺少顶部装饰条")
    if has_header_footer[1]: struct_notes.append("✓ 有底部栏")
    else: struct_score -= 2; struct_notes.append("✗ 缺少底部栏")
    if margin_violations > n: struct_score -= 2; struct_notes.append("✗ 边距不规范")
    else: struct_notes.append("✓ 边距基本规范")
    scores["structure"] = {"score": max(0, struct_score), "max": 15, "notes": struct_notes}
    
    # 2. 排版布局 (15分)
    layout_score = 15
    layout_notes = []
    if overlap_count > 0:
        layout_score -= min(5, overlap_count)
        layout_notes.append(f"✗ 发现{overlap_count}处元素重叠")
    else:
        layout_notes.append("✓ 无元素重叠")
    if empty_count > n * 0.2:
        layout_score -= 3; layout_notes.append(f"✗ 空内容页偏多({empty_count}页)")
    else:
        layout_notes.append("✓ 内容页填充良好")
    # 形状密度
    avg_shapes = shape_count / max(n, 1)
    if avg_shapes < 2: layout_score -= 2; layout_notes.append("✗ 页面元素偏少")
    elif avg_shapes > 15: layout_score -= 2; layout_notes.append("✗ 页面元素过多")
    else: layout_notes.append(f"✓ 元素密度适中({avg_shapes:.1f}个/页)")
    scores["layout"] = {"score": max(0, layout_score), "max": 15, "notes": layout_notes}
    
    # 3. 色彩搭配 (15分)
    color_score = 15
    color_notes = []
    unique_colors = len(set(all_colors))
    if unique_colors < 3: color_score -= 5; color_notes.append("✗ 颜色过于单调")
    elif unique_colors > 15: color_score -= 3; color_notes.append("✗ 颜色过多(>15种)")
    else: color_notes.append(f"✓ 用色适中({unique_colors}种)")
    # 检查是否有强调色（至少有一种亮色）
    bright_colors = [c for c in set(all_colors) if max(int(c[i:i+2],16) for i in (0,2,4)) > 180]
    if bright_colors:
        color_notes.append(f"✓ 有{len(bright_colors)}种亮色点缀")
    else:
        color_score -= 3; color_notes.append("✗ 缺少亮色点缀")
    scores["color"] = {"score": max(0, color_score), "max": 15, "notes": color_notes}
    
    # 4. 字体规范 (15分)
    font_score = 15
    font_notes = []
    unique_fonts = set(all_fonts)
    if len(unique_fonts) > 3: font_score -= 3; font_notes.append(f"✗ 字体过多({len(unique_fonts)}种)")
    elif len(unique_fonts) <= 1: font_score -= 1; font_notes.append("仅用1种字体")
    else: font_notes.append(f"✓ 字体规范({len(unique_fonts)}种)")
    # 字体大小层次
    if all_font_sizes:
        sizes = sorted(set(s/12700 for s in all_font_sizes))
        if len(sizes) < 2: font_score -= 2; font_notes.append("✗ 缺少字号层次")
        elif len(sizes) >= 3: font_notes.append(f"✓ 字号层次丰富({len(sizes)}级)")
        else: font_notes.append(f"✓ 有字号层次({len(sizes)}级)")
    # 是否有大标题
    large_sizes = [s for s in all_font_sizes if s/12700 >= 18]
    if large_sizes: font_notes.append("✓ 有大标题")
    else: font_score -= 2; font_notes.append("✗ 缺少大标题(≥18pt)")
    scores["typography"] = {"score": max(0, font_score), "max": 15, "notes": font_notes}
    
    # 5. 内容呈现 (15分)
    content_score = 15
    content_notes = []
    avg_text_len = total_text_len / max(n, 1)
    if avg_text_len < 50: content_score -= 3; content_notes.append("✗ 内容偏少")
    elif avg_text_len > 500: content_score -= 3; content_notes.append("✗ 文字密集(平均>500字/页)")
    else: content_notes.append(f"✓ 内容密度适中({avg_text_len:.0f}字/页)")
    if image_count > n * 0.15:
        content_notes.append(f"✓ 有配图({image_count}张)")
    else:
        content_score -= 2; content_notes.append("✗ 缺少配图")
    if chart_count > 0:
        content_notes.append(f"✓ 有图表({chart_count}个)")
    else:
        content_score -= 1; content_notes.append("可增加图表")
    scores["content"] = {"score": max(0, content_score), "max": 15, "notes": content_notes}
    
    # 6. 视觉一致性 (10分)
    consist_score = 10
    consist_notes = []
    # 对齐方式一致性
    if alignments:
        most_common = alignments.most_common(1)[0][1]
        total_align = sum(alignments.values())
        if most_common / total_align > 0.7:
            consist_notes.append("✓ 对齐方式统一")
        else:
            consist_score -= 2; consist_notes.append("✗ 对齐方式混乱")
    # 形状数量变化
    if n > 1:
        shapes_per_slide = [len(list(sl.shapes)) for sl in slides]
        max_sh = max(shapes_per_slide)
        min_sh = min(shapes_per_slide)
        if max_sh - min_sh > 10:
            consist_score -= 2; consist_notes.append("✗ 页面复杂度差异大")
        else:
            consist_notes.append("✓ 页面复杂度均匀")
    scores["consistency"] = {"score": max(0, consist_score), "max": 10, "notes": consist_notes}
    
    # 7. 图表数据 (10分)
    chart_score = 10
    chart_notes = []
    if chart_count > 0:
        chart_notes.append(f"✓ 包含{chart_count}个图表")
        # 检查是否有数据标签
        for slide in slides:
            for shape in slide.shapes:
                if hasattr(shape, 'chart'):
                    try:
                        if shape.chart.plots[0].has_data_labels:
                            chart_notes.append("✓ 图表有数据标签")
                        else:
                            chart_score -= 2; chart_notes.append("✗ 部分图表缺少数据标签")
                    except: pass
    else:
        chart_score -= 4; chart_notes.append("✗ 无图表（非必需，但推荐）")
    scores["charts"] = {"score": max(0, chart_score), "max": 10, "notes": chart_notes}
    
    # 8. 比例与尺寸 (5分)
    ratio_score = 5
    ratio_notes = []
    ratio = sw / sh
    if abs(ratio - 16/9) < 0.05:
        ratio_notes.append("✓ 标准16:9比例")
    elif abs(ratio - 4/3) < 0.05:
        ratio_notes.append("4:3比例（建议升级16:9）")
        ratio_score -= 2
    else:
        ratio_notes.append(f"非常规比例{ratio:.2f}")
        ratio_score -= 3
    scores["ratio"] = {"score": max(0, ratio_score), "max": 5, "notes": ratio_notes}
    
    # ════════════════════════════════════
    # 综合评分
    # ════════════════════════════════════
    total = sum(s["score"] for s in scores.values())
    max_total = sum(s["max"] for s in scores.values())
    
    # 等级
    if total >= 90: grade = "🥇 优秀"
    elif total >= 75: grade = "🥈 良好"
    elif total >= 60: grade = "🥉 合格"
    else: grade = "⚠️ 需优化"
    
    report["scores"] = scores
    report["total"] = {"score": total, "max": max_total, "grade": grade}
    report["summary"] = {
        "shapes": shape_count,
        "text_shapes": text_count,
        "charts": chart_count,
        "images": image_count,
        "empty_pages": empty_count,
        "overlaps": overlap_count,
        "fonts": len(unique_fonts),
        "colors": unique_colors,
    }
    
    return report

def print_report(report):
    """打印格式化的评分报告"""
    if "error" in report:
        print(f"❌ {report['error']}")
        return
    
    print(f"\n═══ PPT质量评估报告 ═══")
    print(f"📋 文件: {report['file']}")
    print(f"📄 页数: {report['slide_count']} 页 | 尺寸: {report['slide_size']}")
    print(f"📊 综合得分: {report['total']['score']}/{report['total']['max']} — {report['total']['grade']}\n")
    
    print("━━━ 各维度评分 ━━━\n")
    
    dim_names = {
        "structure": "① 页面结构", "layout": "② 排版布局",
        "color": "③ 色彩搭配", "typography": "④ 字体规范",
        "content": "⑤ 内容呈现", "consistency": "⑥ 视觉一致性",
        "charts": "⑦ 图表数据", "ratio": "⑧ 比例尺寸"
    }
    
    for key, name in dim_names.items():
        s = report["scores"].get(key, {})
        sc = s.get("score", 0)
        mx = s.get("max", 10)
        notes = s.get("notes", [])
        bar = "█" * (sc * 20 // mx) + "░" * (20 - sc * 20 // mx)
        print(f"{name}: [{sc}/{mx}] {bar}")
        for note in notes:
            print(f"   {note}")
        print()
    
    print(f"━━━ 综合评估 ━━━")
    print(f"{report['total']['grade']}  {report['total']['score']}/{report['total']['max']} 分")
    
    # 优化建议
    weak = [(k, s) for k, s in report["scores"].items() 
            if s["score"] / s["max"] < 0.6]
    if weak:
        print(f"\n━━━ 优化建议 ━━━")
        for k, s in weak:
            pct = s["score"] / s["max"]
            dim_name = dim_names.get(k, k)
            print(f"🔴 {dim_name} ({s['score']}/{s['max']})")
            for note in s["notes"]:
                if note.startswith("✗"):
                    print(f"   → {note.replace('✗ ', '')}")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("用法: python3 analyze_ppt.py <path_to_pptx>")
        sys.exit(1)
    
    report = analyze_pptx(sys.argv[1])
    
    if "--json" in sys.argv:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        print_report(report)